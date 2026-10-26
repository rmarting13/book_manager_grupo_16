import csv
import os
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any, Dict, List, Optional, Sequence, Union
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from book_manager.utilities.utilities import PrinterUtil


class ReporteError(ValueError):
    """Clase de error base para los reportes."""
    pass


PATRON_FECHA = re.compile(r"alerta_precios_(\d{8})\.csv$", re.IGNORECASE)

CABECERAS_REPORTE: List[str] = [
    "Libro",
    "Precio interno",
    "Precio web",
    "Diferencia",
    "Fecha de extracción",
]


class LectorAlertas:
    """Lee los archivos CSV de alertas de precios generados por el scraper.

    El nombre del archivo tiene el formato ``alerta_precios_YYYYMMDD.csv`` y la
    fecha de extracción se obtiene de ese nombre.
    """

    def __init__(self, ruta_csv: str) -> None:
        """Inicializa el lector con la ruta del archivo de alertas.

        Args:
            ruta_csv: Ruta del archivo ``alerta_precios_YYYYMMDD.csv``.

        Raises:
            ReporteError: Si el nombre del archivo no contiene una fecha válida.
        """
        self._ruta_csv = ruta_csv
        self._fecha = self._fecha_desde_nombre(ruta_csv)

    @property
    def ruta_csv(self) -> str:
        """Ruta del archivo de alertas leído."""
        return self._ruta_csv

    @property
    def fecha_extraccion(self) -> date:
        """Fecha de extracción derivada del nombre del archivo."""
        return self._fecha

    @staticmethod
    def _fecha_desde_nombre(ruta_csv: str) -> date:
        """Extrae la fecha de extracción del nombre del archivo.

        Args:
            ruta_csv: Ruta del archivo ``alerta_precios_YYYYMMDD.csv``.

        Returns:
            La fecha indicada en el nombre del archivo.

        Raises:
            ReporteError: Si el nombre no respeta el formato esperado.
        """
        coincidencia = PATRON_FECHA.search(os.path.basename(ruta_csv))
        if coincidencia is None:
            raise ReporteError(
                "El nombre del archivo debe tener el formato "
                f"'alerta_precios_AAAAMMDD.csv': {ruta_csv}"
            )
        try:
            return datetime.strptime(coincidencia.group(1), "%Y%m%d").date()
        except ValueError:
            raise ReporteError(
                f"La fecha del nombre del archivo no es válida: {ruta_csv}"
            )

    def leer(self) -> List[Dict[str, Any]]:
        """Lee el CSV y devuelve una fila por libro comparado.

        Returns:
            Lista de diccionarios con el libro, los precios, la diferencia y la
            fecha de extracción.

        Raises:
            ReporteError: Si el archivo no existe o no tiene las columnas
                esperadas.
        """
        if not os.path.exists(self._ruta_csv):
            raise ReporteError(f"No se encontró el archivo de alertas: {self._ruta_csv}")

        with open(self._ruta_csv, encoding="utf-8", newline="") as archivo:
            lector = csv.DictReader(archivo)
            columnas = lector.fieldnames or []
            faltantes = [
                columna
                for columna in ("titulo_catalogo", "titulo_cuspide", "precio_catalogo", "precio_cuspide")
                if columna not in columnas
            ]
            if faltantes:
                raise ReporteError(
                    f"El archivo {self._ruta_csv} no tiene las columnas: "
                    f"{', '.join(faltantes)}"
                )
            return [self._fila(fila) for fila in lector]

    def _fila(self, fila: Dict[str, str]) -> Dict[str, Any]:
        """Convierte una fila del CSV en una fila del reporte.

        Args:
            fila: Fila leída del CSV.

        Returns:
            Diccionario con las columnas del reporte.
        """
        precio_interno = self._a_decimal(fila.get("precio_catalogo"))
        precio_web = self._a_decimal(fila.get("precio_cuspide"))
        diferencia = (
            (precio_web - precio_interno).quantize(Decimal("0.01"))
            if precio_interno is not None and precio_web is not None
            else None
        )
        return {
            "libro": fila.get("titulo_catalogo") or fila.get("titulo_cuspide") or "",
            "precio_interno": precio_interno,
            "precio_web": precio_web,
            "diferencia": diferencia,
            "fecha_extraccion": self._fecha,
        }

    @staticmethod
    def _a_decimal(valor: Optional[str]) -> Optional[Decimal]:
        """Convierte un precio del CSV a Decimal, o None si está vacío.

        Args:
            valor: Texto leído del CSV.

        Returns:
            El precio como Decimal, o ``None`` si la celda está vacía o no es
            numérica.
        """
        if valor is None or valor.strip() == "":
            return None
        try:
            return Decimal(valor.strip())
        except (InvalidOperation, ValueError):
            return None


class ReporteAlertasPrecios:
    """Genera un reporte Excel con los precios del catálogo y del sitio web.

    Lee uno o varios archivos ``alerta_precios_YYYYMMDD.csv`` y escribe un
    ``.xlsx`` con las columnas Libro, Precio interno, Precio web, Diferencia y
    Fecha de extracción. La fecha de cada fila se deriva del nombre del archivo
    de origen, por eso un mismo libro puede aparecer una vez por cada día en
    que se ejecutó el scraper.
    """

    def __init__(self, rutas_csv: Union[str, Sequence[str]]) -> None:
        """Inicializa el reporte con uno o varios archivos de alertas.

        Args:
            rutas_csv: Ruta de un archivo ``alerta_precios_YYYYMMDD.csv`` o una
                secuencia con las rutas de varios archivos.

        Raises:
            ReporteError: Si no se recibe ningún archivo.
        """
        if isinstance(rutas_csv, str):
            rutas_csv = [rutas_csv]
        if not rutas_csv:
            raise ReporteError("Se necesita al menos un archivo de alertas de precios.")
        self._lectores = [LectorAlertas(ruta) for ruta in rutas_csv]
        self._filas: List[Dict[str, Any]] = []

    @property
    def lectores(self) -> List[LectorAlertas]:
        """Lectores de los archivos de alertas de origen."""
        return self._lectores

    @property
    def fechas_extraccion(self) -> List[date]:
        """Fechas de extracción de los archivos de origen, sin repetir y ordenadas."""
        return sorted({lector.fecha_extraccion for lector in self._lectores})

    @property
    def filas(self) -> List[Dict[str, Any]]:
        """Filas del reporte leídas de todos los archivos de origen."""
        return self._filas

    def generar(self, ruta_xlsx: Optional[str] = None) -> str:
        """Genera el archivo Excel con las alertas de precios.

        Args:
            ruta_xlsx: Ruta del archivo a crear. Si es ``None``, se usa
                ``reporte_alertas_precios.xlsx`` junto al primer CSV.

        Returns:
            La ruta del archivo Excel generado.

        Raises:
            ReporteError: Si alguno de los CSV de origen no se puede leer.
        """
        self._filas = self._leer_todos()
        ruta = ruta_xlsx or self.ruta_por_defecto()

        libro = Workbook()
        hoja = libro.active
        assert isinstance(hoja, Worksheet)
        hoja.title = "Alertas de precios"
        self._escribir_cabecera(hoja)
        for numero, fila in enumerate(self._filas, start=2):
            self._escribir_fila(hoja, numero, fila)
        self._ajustar_anchos(hoja)

        carpeta = os.path.dirname(os.path.abspath(ruta))
        os.makedirs(carpeta, exist_ok=True)
        libro.save(ruta)
        return ruta

    def ruta_por_defecto(self) -> str:
        """Ruta sugerida del Excel, junto al primer CSV de origen."""
        carpeta = os.path.dirname(os.path.abspath(self._lectores[0].ruta_csv))
        return os.path.join(carpeta, "reporte_alertas_precios.xlsx")

    def _leer_todos(self) -> List[Dict[str, Any]]:
        """Lee todos los archivos de alertas y junta sus filas.

        Returns:
            Las filas de todos los archivos, ordenadas por fecha de extracción
            y, dentro de cada fecha, por libro.
        """
        filas: List[Dict[str, Any]] = []
        for lector in self._lectores:
            filas.extend(lector.leer())
        return sorted(filas, key=lambda fila: (fila["fecha_extraccion"], fila["libro"]))

    @staticmethod
    def _escribir_cabecera(hoja: Worksheet) -> None:
        """Escribe y da formato a la fila de cabeceras.

        Args:
            hoja: Hoja de cálculo donde se escribe el reporte.
        """
        relleno = PatternFill("solid", fgColor="1F4E78")
        for columna, titulo in enumerate(CABECERAS_REPORTE, start=1):
            celda = hoja.cell(row=1, column=columna, value=titulo)
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = relleno
            celda.alignment = Alignment(horizontal="center")

    @staticmethod
    def _escribir_fila(hoja: Worksheet, numero: int, fila: Dict[str, Any]) -> None:
        """Escribe una fila del reporte con los formatos de precio y fecha.

        Args:
            hoja: Hoja de cálculo donde se escribe el reporte.
            numero: Número de fila de Excel donde escribir.
            fila: Fila del reporte a escribir.
        """
        valores = [
            fila["libro"],
            ReporteAlertasPrecios._a_float(fila["precio_interno"]),
            ReporteAlertasPrecios._a_float(fila["precio_web"]),
            ReporteAlertasPrecios._a_float(fila["diferencia"]),
            fila["fecha_extraccion"],
        ]
        for columna, valor in enumerate(valores, start=1):
            celda = hoja.cell(row=numero, column=columna, value=valor)
            if columna in (2, 3, 4):
                celda.number_format = '#,##0.00'
            if columna == 5:
                celda.number_format = 'DD/MM/YYYY'

    @staticmethod
    def _ajustar_anchos(hoja: Worksheet) -> None:
        """Ajusta el ancho de cada columna al contenido.

        Args:
            hoja: Hoja de cálculo a ajustar.
        """
        for columna, titulo in enumerate(CABECERAS_REPORTE, start=1):
            ancho = len(titulo) + 4
            for celda in hoja[get_column_letter(columna)]:
                if celda.value is not None:
                    ancho = max(ancho, len(str(celda.value)) + 4)
            hoja.column_dimensions[get_column_letter(columna)].width = ancho

    @staticmethod
    def _a_float(valor: Optional[Decimal]) -> Optional[float]:
        """Convierte un Decimal a float para escribirlo en la planilla.

        Args:
            valor: Valor a convertir.

        Returns:
            El valor como float, o ``None`` si no hay precio.
        """
        return float(valor) if valor is not None else None


# Pruebas unitarias:
if __name__ == '__main__':
    import tempfile

    # 1. Fecha derivada del nombre del archivo
    PrinterUtil.dibujar_separador("1. Fecha derivada del nombre del archivo")

    lector = LectorAlertas("downloads/alerta_precios_20261026.csv")
    print(f"Archivo : {lector.ruta_csv}")
    print(f"Fecha   : {lector.fecha_extraccion}")

    try:
        LectorAlertas("downloads/alertas_precios.csv")
    except ReporteError as error:
        print(f"Error esperado: {error}")

    # 2. Lectura del CSV real del scraper
    PrinterUtil.dibujar_separador("2. Lectura del CSV de alertas")

    ruta_csv_real = os.path.join(
        os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
        "scraper",
        "downloads",
        "alerta_precios_20261026.csv",
    )
    if os.path.exists(ruta_csv_real):
        for fila in LectorAlertas(ruta_csv_real).leer():
            print(
                f"{fila['libro']:<28} | interno: {fila['precio_interno']} | "
                f"web: {fila['precio_web']} | dif: {fila['diferencia']} | "
                f"fecha: {fila['fecha_extraccion']}"
            )
    else:
        print(f"No se encontró el CSV real: {ruta_csv_real}")

    # 3. Generación del Excel a partir de varios CSV de prueba
    PrinterUtil.dibujar_separador("3. Generación del Excel con varios CSV")
    
    carpeta_temporal = "src/book_manager/reportes"

    def escribir_csv(fecha: str, filas: List[List[str]]) -> str:
        """Escribe un CSV de alertas de prueba con el nombre del scraper."""
        ruta = os.path.join(carpeta_temporal, f"alerta_precios_{fecha}.csv")
        with open(ruta, "w", encoding="utf-8", newline="") as archivo:
            escritor = csv.writer(archivo)
            escritor.writerow(
                [
                    "isbn_catalogo", "isbn_cuspide", "titulo_catalogo", "titulo_cuspide",
                    "autor_catalogo", "autor_cuspide", "coincidencia", "moneda",
                    "precio_catalogo", "precio_cuspide", "diferencia",
                    "variacion_porcentual", "estado", "url",
                ]
            )
            escritor.writerows(filas)
        return ruta

    ruta_csv_ayer = escribir_csv(
        "20261025",
        [
            ["9788420414706", "9789877252538", "Rayuela", "RAYUELA", "Julio Cortázar",
             "CORTAZAR, JULIO", "TITULO", "$", "15900.00", "45999.00", "30099.00",
             "189.30", "ALERTA_SUBE", "https://cuspide.com/producto/rayuela/"],
            ["9789873650383", "9789873650383", "Maus", "MAUS", "Art Spiegelman",
             "SPIEGELMAN, ART", "TITULO", "$", "", "56999.00", "", "",
             "SIN_PRECIO_CATALOGO_CON_USD", "https://cuspide.com/producto/maus/"],
        ],
    )
    ruta_csv_hoy = escribir_csv(
        "20261026",
        [
            ["9788420414706", "9789877252538", "Rayuela", "RAYUELA", "Julio Cortázar",
             "CORTAZAR, JULIO", "TITULO", "$", "15900.00", "47499.00", "31599.00",
             "198.74", "ALERTA_SUBE", "https://cuspide.com/producto/rayuela/"],
            ["9788439732471", "9789877256963", "Cien años de soledad",
             "CIEN AÑOS DE SOLEDAD", "Gabriel García Márquez", "GARCIA MARQUEZ, GABRIEL",
             "TITULO", "$", "14200.00", "38999.00", "24799.00", "174.64",
             "ALERTA_SUBE", "https://cuspide.com/producto/cien-anos-de-soledad/"],
        ],
    )

    reporte = ReporteAlertasPrecios([ruta_csv_hoy, ruta_csv_ayer])
    ruta_xlsx = reporte.generar(os.path.join(carpeta_temporal, "reporte_varios.xlsx"))
    print(f"Excel generado : {ruta_xlsx}")
    print(f"Archivos leídos: {len(reporte.lectores)}")
    print(f"Fechas         : {reporte.fechas_extraccion}")
    print(f"Filas          : {len(reporte.filas)}")

    # 5. Un solo archivo también es válido
    PrinterUtil.dibujar_separador("5. Un único archivo")

    reporte_unico = ReporteAlertasPrecios(ruta_csv_hoy)
    ruta_xlsx_unico = reporte_unico.generar()
    print(f"Excel generado : {ruta_xlsx_unico}")
    print(f"Filas          : {len(reporte_unico.filas)}")
    assert os.path.basename(ruta_xlsx_unico) == "reporte_alertas_precios.xlsx"

    try:
        ReporteAlertasPrecios([])
    except ReporteError as error:
        print(f"Error esperado: {error}")